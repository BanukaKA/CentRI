"""Generate the concise Render-ready CentRI presentation.

Usage:
    .venv\\Scripts\\python scripts/generate_short_presentation.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from scripts.generate_presentation import (
    BLUE,
    BLUE_LIGHT,
    GREEN,
    GREEN_LIGHT,
    INK,
    MUTED,
    NAVY,
    NAVY_2,
    PAPER,
    PURPLE,
    PURPLE_LIGHT,
    RED,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    arrow,
    card,
    footer,
    set_bg,
    shape,
    textbox,
    title,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "CentRI_concise_presentation.pptx"


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1: thesis.
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.65, 0, 3.68, 7.5, NAVY_2, NAVY_2)
    textbox(slide, "CENTRI / CENTRAL ROUTING INTELLIGENCE", 0.75, 0.72, 7.6, 0.25, 11, BLUE_LIGHT, True)
    textbox(slide, "One front door.\nThousands of agents.", 0.72, 1.55, 7.8, 1.25, 32, WHITE, True)
    textbox(
        slide,
        "A semantic API gateway that compiles agent capability offline, "
        "retrieves locally, and uses NVIDIA Nemotron only to resolve ambiguity.",
        0.76,
        3.25,
        7.8,
        0.85,
        18,
        BLUE_LIGHT,
    )
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 10.15, 1.5, 2.7, 0.7, BLUE, BLUE, True)
    textbox(slide, "USER REQUEST", 10.15, 1.73, 2.7, 0.2, 12, WHITE, True, align=PP_ALIGN.CENTER)
    arrow(slide, 11.5, 2.25, 11.5, 2.9, BLUE_LIGHT, 2)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 10.15, 2.95, 2.7, 0.9, TEAL, TEAL, True)
    textbox(slide, "CENTRI\nROUTING PLANE", 10.15, 3.16, 2.7, 0.45, 13, WHITE, True, align=PP_ALIGN.CENTER)
    arrow(slide, 11.5, 3.95, 11.5, 4.6, BLUE_LIGHT, 2)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 10.15, 4.65, 2.7, 0.7, PURPLE, PURPLE, True)
    textbox(slide, "CORRECT AGENT ID", 10.15, 4.88, 2.7, 0.2, 12, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, "2:50 live pitch | Enterprise AI needs a control plane for agent discovery.", 0.78, 6.2, 8.7, 0.3, 14, WHITE, True)
    footer(slide, 1, "Nebius x NVIDIA Global AI Hackathon", dark=True)

    # Slide 2: architecture.
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PAPER)
    title(slide, "0:20-0:40 | THE ARCHITECTURE", "Compile once. Retrieve locally. Reason narrowly.", "Three planes keep the runtime decision surface bounded.")
    card(slide, 0.72, 1.85, 3.65, 3.55, "1  OFFLINE COMPILATION", "Agent schemas\n→ synthetic queries\n→ contrastive hard negatives\n→ Qwen embeddings\n→ HNSW / FAISS routing memory", BLUE_LIGHT, BLUE, 15)
    card(slide, 4.84, 1.85, 3.65, 3.55, "2  RUNTIME ROUTING", "FastAPI gateway\n→ intent shards\n→ local Top-K retrieval\n→ compact routing cards\n→ NVIDIA Nemotron reranking", TEAL_LIGHT, TEAL, 15)
    card(slide, 8.96, 1.85, 3.65, 3.55, "3  GOVERNED HANDOFF", "agent_id\n→ HTTP runtime\n→ concurrent execution\n→ approval gates\n→ redacted telemetry\n→ correction vectors", PURPLE_LIGHT, PURPLE, 15)
    arrow(slide, 4.42, 3.6, 4.78, 3.6, BLUE, 2)
    arrow(slide, 8.54, 3.6, 8.9, 3.6, TEAL, 2)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.05, 5.85, 11.15, 0.65, NAVY, NAVY, True)
    textbox(slide, "Local routing memory + bounded GPU reasoning = scalable semantic discovery", 1.25, 6.06, 10.75, 0.2, 16, WHITE, True, align=PP_ALIGN.CENTER)
    footer(slide, 2, "FastAPI · HNSW/FAISS · Qwen Embeddings · NVIDIA Nemotron · Nebius · Tavily")

    # Slide 3: live demo.
    slide = prs.slides.add_slide(blank)
    set_bg(slide, WHITE)
    title(slide, "0:40-1:05 | LIVE RENDER DEMO", "One public URL, one stable routing contract", "Postman-ready requests are included in postman_test.txt.")
    textbox(slide, "https://centri-demo.onrender.com", 0.85, 1.72, 11.7, 0.5, 24, BLUE, True, align=PP_ALIGN.CENTER)
    steps = [
        ("HEALTH", "/health", "index loaded\n+ test cURLs"),
        ("ROUTE", "/route", "agent_id\n+ candidates"),
        ("CHAT", "/chat", "agent response\n+ approval"),
        ("TAVILY", "use_web_search", "live sources\n+ web_research"),
    ]
    xs = [0.72, 3.78, 6.84, 9.9]
    for index, (label, path, body) in enumerate(steps):
        card(slide, xs[index], 2.75, 2.55, 1.85, label, f"{path}\n\n{body}", [BLUE_LIGHT, TEAL_LIGHT, PURPLE_LIGHT, GREEN_LIGHT][index], [BLUE, TEAL, PURPLE, GREEN][index], 13)
        if index < 3:
            arrow(slide, xs[index] + 2.62, 3.68, xs[index] + 2.92, 3.68, MUTED, 1.5)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.05, 5.25, 11.15, 0.9, NAVY, NAVY, True)
    textbox(slide, "Demo message: CentRI returns the destination agent ID; downstream execution is governed separately.", 1.35, 5.55, 10.55, 0.3, 15, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, "Free Render hosting may cold-start. Call /health first.", 1.2, 6.45, 10.9, 0.2, 11, RED, True, align=PP_ALIGN.CENTER)
    footer(slide, 3, "Live demo: centri-demo.onrender.com")

    # Slide 4: how to test the live deployment.
    slide = prs.slides.add_slide(blank)
    set_bg(slide, WHITE)
    title(slide, "1:05-1:35 | HOW TO TEST", "Use the live Render deployment", "Postman-ready requests are in postman_test.txt.")
    textbox(slide, "https://centri-demo.onrender.com", 0.85, 1.7, 11.7, 0.42, 22, BLUE, True, align=PP_ALIGN.CENTER)
    test_steps = [
        ("1", "Wake it", "Run GET /health first.\nFree Render services may cold-start for about a minute."),
        ("2", "Import", "Import postman_test.txt into Postman.\nUse the Render URL, not localhost."),
        ("3", "Route", "POST /route\nReceive agent_id, candidates, confidence."),
        ("4", "Execute", "POST /chat\nSet approval_granted=true for the demo."),
        ("5", "Research", "Set use_web_search=true\nVerify web_research.status = ok."),
    ]
    for index, (number, heading, body) in enumerate(test_steps):
        y = 2.55 + index * 0.68
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 0.95, y + 0.02, 0.38, 0.38, [BLUE, TEAL, PURPLE, GREEN, NAVY][index], [BLUE, TEAL, PURPLE, GREEN, NAVY][index])
        textbox(slide, number, 0.95, y + 0.11, 0.38, 0.15, 10, WHITE, True, align=PP_ALIGN.CENTER)
        textbox(slide, heading, 1.55, y, 1.35, 0.25, 14, [BLUE, TEAL, PURPLE, GREEN, NAVY][index], True)
        textbox(slide, body, 3.0, y, 8.9, 0.4, 13, INK)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.15, 6.05, 11.0, 0.55, TEAL_LIGHT, TEAL, True)
    textbox(slide, "Expected sample behavior: agent execution may be marked simulated because no public agent runtime is configured.", 1.4, 6.23, 10.5, 0.18, 11.5, TEAL, True, align=PP_ALIGN.CENTER)
    footer(slide, 4, "Live demo: centri-demo.onrender.com")

    # Slide 5: scale and evidence.
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    title(slide, "1:35-2:25 | SCALE + EVIDENCE", "The model should not see the whole enterprise catalog", "CentRI separates catalog size from runtime reasoning context.", dark=True)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.75, 1.82, 5.65, 4.4, NAVY_2, BLUE, True)
    textbox(slide, "PROMPT-ALL ROUTING", 1.05, 2.12, 2.8, 0.25, 12, RED, True)
    textbox(slide, "C ≈ A × schema_tokens", 1.05, 2.72, 4.9, 0.5, 24, WHITE, True)
    textbox(slide, "Every new agent expands the request-time context, cost, and ambiguity.", 1.05, 3.48, 4.8, 0.65, 16, BLUE_LIGHT)
    textbox(slide, "CENTRI ROUTING", 1.05, 4.52, 2.8, 0.25, 12, TEAL_LIGHT, True)
    textbox(slide, "C ≈ query + K × routing_card", 1.05, 5.08, 5.0, 0.5, 22, WHITE, True)
    textbox(slide, "K << A  |  local HNSW discovery  |  bounded Nemotron GPU context", 1.05, 5.72, 5.0, 0.3, 13, TEAL_LIGHT, True)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.72, 1.82, 5.85, 4.4, TEAL, TEAL, True)
    textbox(slide, "CAPTURED 20-CASE BENCHMARK", 7.05, 2.12, 4.7, 0.25, 12, WHITE, True)
    textbox(slide, "100%", 7.05, 2.75, 1.35, 0.55, 30, WHITE, True)
    textbox(slide, "candidate recall", 7.05, 3.28, 2.2, 0.25, 13, TEAL_LIGHT)
    textbox(slide, "100%", 9.55, 2.75, 1.35, 0.55, 30, WHITE, True)
    textbox(slide, "selective accuracy", 9.55, 3.28, 2.35, 0.25, 13, TEAL_LIGHT)
    textbox(slide, "80% coverage  |  3.09s mean  |  5.53s p95", 7.05, 4.35, 4.85, 0.3, 16, WHITE, True)
    textbox(slide, "Uncertain cases abstain instead of creating an incorrect handoff.", 7.05, 5.12, 4.75, 0.65, 16, WHITE)
    footer(slide, 5, "Measured on synthetic golden cases; larger-scale load testing is next", dark=True)

    # Slide 6: close.
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PAPER)
    title(slide, "2:25-2:50 | THE TAKEAWAY", "CentRI turns an agent network into a semantic API", "A shared enterprise capability layer with one efficient entry point.")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.85, 1.9, 11.65, 1.0, NAVY, NAVY, True)
    textbox(slide, "user request  ->  CentRI  ->  agent_id  ->  authorized runtime", 1.1, 2.25, 11.1, 0.3, 20, WHITE, True, align=PP_ALIGN.CENTER)
    bullet_items = [
        ("Speed", "local semantic retrieval"),
        ("Accuracy", "bounded Nemotron reranking"),
        ("Safety", "abstention and approval gates"),
        ("Scale", "compiled memory and HNSW"),
        ("Access", "approved agents and enterprise resources for authorized users"),
    ]
    for index, (label, body) in enumerate(bullet_items):
        y = 3.45 + index * 0.5
        textbox(slide, label, 1.25, y, 1.3, 0.22, 14, [BLUE, TEAL, PURPLE, GREEN, NAVY][index], True)
        textbox(slide, body, 2.75, y, 6.9, 0.22, 14, INK)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.7, 3.5, 2.25, 1.6, TEAL_LIGHT, TEAL, True)
    textbox(slide, "Render demo\n+ Postman\n+ Tavily live", 9.7, 3.93, 2.25, 0.72, 16, TEAL, True, align=PP_ALIGN.CENTER)
    textbox(slide, "https://centri-demo.onrender.com/docs", 1.25, 6.35, 10.8, 0.25, 15, BLUE, True, align=PP_ALIGN.CENTER)
    footer(slide, 6, "CentRI / Central Routing Intelligence")

    return prs


def main() -> None:
    build().save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
